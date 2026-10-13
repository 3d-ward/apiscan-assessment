"""Spreadsheet renderer.

EXTRACTED CAPABILITY (CAP-032). Four sheets, matching the source's structure
because it proved to be what a client actually reads:

  Summary       engagement metadata, severity counts, execution status
  Findings      the register, one row per distinct defect
  Test Log      every case with its disposition, so coverage is auditable
  Coverage      the honesty block: what was not established

``openpyxl`` is an OPTIONAL dependency. It is imported inside ``write`` rather
than at module scope so that importing the reporting package never fails, and
so the absence of the package downgrades the deliverable to Markdown instead of
breaking the run. ``available()`` lets a caller check first.

Image embedding needs ``Pillow`` in addition. That is reported honestly rather
than silently skipped: ``write`` returns a result object naming what was and
was not included.
"""
from . import safe

SEVERITY_FILL = {
    "Critical": "7B1E1E", "High": "C0392B", "Medium": "E67E22",
    "Low": "F1C40F", "Informational": "5DADE2",
}
SEVERITY_FONT = {
    "Critical": "FFFFFF", "High": "FFFFFF", "Medium": "FFFFFF",
    "Low": "000000", "Informational": "000000",
}
KIND_FILL = {
    "DEFECT": "F8CBAD", "UNRESOLVED": "FFE699",
    "CONTROL": "C6EFCE", "SCOPE-EXCLUDED": "D9D9D9",
}
VERDICT_FILL = {
    "CONTROL-PRESENT": "C6EFCE", "CONTROL-ABSENT": "FFC7CE",
    "OVER-BLOCKED": "FFE699", "BLOCKED": "D9D9D9",
    "ERROR": "F4CCCC", "OBSERVED": "DEEAF6", "UNCLEAR": "FFF2CC",
    "SKIPPED": "EEEEEE",
}
HEADER_FILL = "1F3864"


class RenderResult(object):
    def __init__(self, path, sheets, images_embedded, notes):
        self.path = path
        self.sheets = sheets
        self.images_embedded = images_embedded
        self.notes = list(notes)

    def __repr__(self):
        return "<RenderResult %s sheets=%d images=%d>" % (
            self.path, len(self.sheets), self.images_embedded)


def available():
    try:
        import openpyxl           # noqa: F401
        return True
    except ImportError:
        return False


def pillow_available():
    try:
        import PIL                # noqa: F401
        return True
    except ImportError:
        return False


def write(model, path, poc_images=None):
    """Write the workbook. Raises DependencyUnavailable if openpyxl is absent."""
    from ...core.errors import DependencyUnavailable
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise DependencyUnavailable(
            "openpyxl", "spreadsheet report rendering",
            "python3 -m pip install openpyxl")

    notes = []
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="top")
    top = Alignment(vertical="top")
    hdr_font = Font(bold=True, color="FFFFFF", size=10)
    hdr_fill = PatternFill("solid", fgColor=HEADER_FILL)
    title_font = Font(bold=True, size=14, color="1F3864")

    wb = openpyxl.Workbook()
    sheets = []

    def header(ws, row, labels, widths):
        for i, label in enumerate(labels, 1):
            cell = ws.cell(row=row, column=i, value=safe.cell(label))
            cell.font = hdr_font
            cell.fill = hdr_fill
            cell.border = border
            cell.alignment = wrap
            ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]
        ws.row_dimensions[row].height = 26
        ws.freeze_panes = ws.cell(row=row + 1, column=1)

    # -- Summary ---------------------------------------------------------
    ws = wb.active
    ws.title = "Summary"
    sheets.append(ws.title)
    sm = model.summary()
    ws["A1"] = "%s -- API security assessment" % sm["engagement"]
    ws["A1"].font = title_font
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 62
    rows = [
        ("Target", sm["target"]),
        ("Generated", sm["generated"]),
        ("Run id", sm["run_id"] or "n/a"),
        ("Run status", sm["status"]),
        ("", ""),
        ("Confirmed defects", sm["confirmed_defects"]),
        ("Unresolved (impact not externally provable)", sm["unresolved"]),
        ("Controls verified as present", sm["verified_controls"]),
        ("Scope-excluded items", sm["scope_excluded"]),
        ("", ""),
        ("Cases planned", sm["tally"]["planned"]),
        ("Cases executed", sm["tally"]["executed"]),
        ("Skipped (out of scope)", sm["tally"]["skipped_out_of_scope"]),
        ("Blocked before the application", sm["tally"]["blocked_before_application"]),
        ("Errored", sm["tally"]["errored"]),
        ("Unclear", sm["tally"]["unclear"]),
        ("Conclusive cases", sm["tally"]["conclusive"]),
        ("Execution coverage %", sm["percentages"]["execution_coverage_pct"]),
        ("Conclusive proportion of executed %", sm["percentages"]["conclusive_pct"]),
        ("", ""),
        ("NOTE", "Blocked and skipped cases are listed separately and are not "
                 "counted as controls present. An unreached test is not "
                 "evidence about the control it was intended to probe."),
    ]
    r = 3
    for label, value in rows:
        ws.cell(row=r, column=1, value=safe.cell(label)).font = Font(bold=bool(label))
        c = ws.cell(row=r, column=2, value=safe.cell(value))
        c.alignment = wrap
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="Confirmed defects by severity").font = \
        Font(bold=True)
    r += 1
    for sev in ("Critical", "High", "Medium", "Low", "Informational"):
        ws.cell(row=r, column=1, value=sev)
        c = ws.cell(row=r, column=2, value=sm["by_severity"].get(sev, 0))
        fill = SEVERITY_FILL.get(sev)
        if fill:
            ws.cell(row=r, column=1).fill = PatternFill("solid", fgColor=fill)
            ws.cell(row=r, column=1).font = Font(
                bold=True, color=SEVERITY_FONT.get(sev, "000000"))
        c.alignment = top
        r += 1

    # -- Findings --------------------------------------------------------
    ws = wb.create_sheet("Findings")
    sheets.append(ws.title)
    labels = ["ID", "Kind", "Severity", "Title", "Category", "CWE", "OWASP",
              "Affected endpoints", "Endpoints", "Occurrences",
              "Supporting cases", "Provenance", "Capture method",
              "Description", "Impact", "Recommendation", "Conclusion",
              "Evidence", "Notes"]
    widths = [9, 15, 13, 42, 20, 30, 16, 28, 10, 10, 26, 18, 16, 60, 60, 60,
              40, 34, 30]
    header(ws, 1, labels, widths)
    r = 2
    for row in model.findings_rows():
        values = [
            row["id"], row["kind"], row["severity"], row["title"],
            row["category"], row["cwe"], row["owasp"], row["affected"],
            row["endpoint_count"], row["occurrences"], row["cases"],
            row["provenance"], ", ".join(row["capture_methods"]),
            row["description"], row["impact"], row["recommendation"],
            row["conclusion"], "\n".join(row["evidence_files"]), row["notes"],
        ]
        for i, value in enumerate(values, 1):
            cell = ws.cell(row=r, column=i, value=safe.cell(value))
            cell.border = border
            cell.alignment = wrap
        kf = KIND_FILL.get(row["kind"])
        if kf:
            ws.cell(row=r, column=2).fill = PatternFill("solid", fgColor=kf)
        sf = SEVERITY_FILL.get(row["severity"])
        if sf:
            c = ws.cell(row=r, column=3)
            c.fill = PatternFill("solid", fgColor=sf)
            c.font = Font(bold=True,
                          color=SEVERITY_FONT.get(row["severity"], "000000"))
        r += 1

    # -- Test Log --------------------------------------------------------
    ws = wb.create_sheet("Test Log")
    sheets.append(ws.title)
    labels = ["Case", "Phase", "Family", "Endpoint", "Method", "Payload class",
              "Expectation", "Executed", "HTTP", "Layer", "Verdict",
              "Control depth", "Capture method", "Reason", "Evidence",
              "Description"]
    widths = [14, 9, 9, 20, 9, 15, 13, 10, 8, 20, 18, 14, 16, 54, 34, 70]
    header(ws, 1, labels, widths)
    r = 2
    for row in model.test_log_rows():
        values = [row["case_id"], row["phase"], row["family"], row["endpoint"],
                  row["method"], row["payload_class"], row["expectation"],
                  row["executed"], row["http_status"], row["layer"],
                  row["verdict"], row["control_depth"], row["capture_method"],
                  row["reason"], row["evidence"], row["description"]]
        for i, value in enumerate(values, 1):
            cell = ws.cell(row=r, column=i, value=safe.cell(value))
            cell.border = border
            cell.alignment = wrap
        vf = VERDICT_FILL.get(row["verdict"])
        if vf:
            ws.cell(row=r, column=11).fill = PatternFill("solid", fgColor=vf)
        r += 1

    # -- Coverage --------------------------------------------------------
    ws = wb.create_sheet("Coverage")
    sheets.append(ws.title)
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 96
    ws["A1"] = "Coverage and limitations"
    ws["A1"].font = title_font
    cov = model.coverage.statement()
    r = 3
    ws.cell(row=r, column=1, value="Run status").font = Font(bold=True)
    ws.cell(row=r, column=2, value=cov["status"])
    r += 1
    ws.cell(row=r, column=1, value="Baseline calibration").font = Font(bold=True)
    ws.cell(row=r, column=2,
            value="%s -- %s" % ("PASSED" if cov["baseline"]["ok"]
                                else "NOT PASSED",
                                cov["baseline"]["summary"] or "n/a")
            ).alignment = wrap
    r += 2
    ws.cell(row=r, column=1,
            value="What this assessment did NOT establish").font = Font(bold=True)
    r += 1
    for item in cov["not_established"]:
        ws.cell(row=r, column=2, value=safe.cell(item)).alignment = wrap
        r += 1
    if cov["problems"]:
        r += 1
        ws.cell(row=r, column=1,
                value="Outstanding issues with this run").font = Font(bold=True)
        r += 1
        for item in cov["problems"]:
            ws.cell(row=r, column=2, value=safe.cell(item)).alignment = wrap
            r += 1
    r += 1
    ws.cell(row=r, column=1, value="Scope note").font = Font(bold=True)
    ws.cell(row=r, column=2, value=model.scope_note or "Not stated"
            ).alignment = wrap

    # -- PoC Evidence (optional, needs Pillow) ---------------------------
    images = 0
    if poc_images:
        if not pillow_available():
            notes.append(
                "PoC images were supplied but Pillow is not installed, so they "
                "were NOT embedded. The workbook references their paths "
                "instead. Install with: python3 -m pip install Pillow")
            ws = wb.create_sheet("PoC Evidence")
            sheets.append(ws.title)
            ws.column_dimensions["A"].width = 20
            ws.column_dimensions["B"].width = 90
            ws["A1"] = "PoC evidence (paths only -- images not embedded)"
            ws["A1"].font = title_font
            r = 3
            for item in poc_images:
                ws.cell(row=r, column=1, value=item.get("finding_id", ""))
                ws.cell(row=r, column=2, value=item.get("path", "")
                        ).alignment = wrap
                ws.cell(row=r + 1, column=2,
                        value="capture method: %s"
                              % item.get("capture_method", "unspecified")
                        ).alignment = wrap
                r += 3
        else:
            from openpyxl.drawing.image import Image as XLImage
            ws = wb.create_sheet("PoC Evidence")
            sheets.append(ws.title)
            ws.column_dimensions["A"].width = 24
            ws["A1"] = "PoC evidence"
            ws["A1"].font = title_font
            r = 3
            for item in poc_images:
                ws.cell(row=r, column=1,
                        value="%s -- %s" % (item.get("finding_id", ""),
                                            item.get("title", ""))
                        ).font = Font(bold=True)
                r += 1
                ws.cell(row=r, column=1,
                        value="capture method: %s"
                              % item.get("capture_method", "unspecified"))
                r += 1
                try:
                    img = XLImage(item["path"])
                    ws.add_image(img, "A%d" % r)
                    images += 1
                    r += 32
                except Exception as exc:
                    ws.cell(row=r, column=1,
                            value="could not embed %s: %s"
                                  % (item.get("path"), exc))
                    r += 2

    wb.save(path)
    return RenderResult(path, sheets, images, notes)
